"""LINEから貼り付けたシフトを、Excelの氏名×日付の表に対応付ける。"""
import base64
import calendar
from datetime import date, datetime
from io import BytesIO
import re
from zipfile import ZipFile

from openpyxl import load_workbook, Workbook
from openpyxl.utils import column_index_from_string


class InputError(ValueError):
    pass


def read_book(encoded):
    try:
        raw = base64.b64decode(encoded, validate=True)
        if len(raw) > 8 * 1024 * 1024:
            raise InputError("Excelは8MB以下にしてください。")
        with ZipFile(BytesIO(raw)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 64 * 1024 * 1024:
                raise InputError("展開後のExcelが大きすぎます。")
            if any(i.filename.endswith('vbaProject.bin') for i in archive.infolist()):
                raise InputError("マクロ付きExcelは未対応です。.xlsx形式のコピーを用意してください。")
        return load_workbook(BytesIO(raw))
    except InputError:
        raise
    except Exception as exc:
        raise InputError("Excelを読めません。.xlsx形式のファイルを選んでください。") from exc


def normalize_name(value):
    return re.sub(r"\s+", "", str(value or ""))


def normalize_numeric(value):
    """日付・時刻の数字と記号だけを変換する。氏名には適用しない。"""
    if not isinstance(value, str):
        return value
    return value.translate(str.maketrans("０１２３４５６７８９：／－", "0123456789:/-"))


def normalize_clock(value):
    value = normalize_numeric(value)
    if not isinstance(value, str):
        return value
    value = value.strip()
    if re.fullmatch(r"[0-9]{1,2}:[0-9]{2}", value):
        hour, minute = map(int, value.split(":"))
        if 0 <= hour <= 23 and 0 <= minute <= 59:
            return f"{hour:02}:{minute:02}"
    return value


def month_parts(value):
    try:
        value = normalize_numeric(value)
        if not re.fullmatch(r"\d{4}-\d{2}", value):
            raise ValueError()
        year, month = map(int, value.split("-"))
        date(year, month, 1)
        return year, month
    except (ValueError, TypeError):
        raise InputError("対象月を選んでください。")


def header_date(value, year, month):
    if isinstance(value, (datetime, date)):
        return value.date() if isinstance(value, datetime) else value
    text = normalize_numeric(str(value or "")).strip()
    try:
        if re.fullmatch(r"\d{1,2}(?:日)?", text):
            return date(year, month, int(text.rstrip("日")))
        if re.fullmatch(r"\d{1,2}/\d{1,2}", text):
            m, d = map(int, text.split("/"))
            return date(year, m, d)
        if re.fullmatch(r"\d{4}[-/]\d{1,2}[-/]\d{1,2}", text):
            return date(*map(int, re.split(r"[-/]", text)))
    except ValueError:
        pass
    return None


def parse_lines(text, month_value):
    year, month = month_parts(month_value)
    if not text.strip():
        raise InputError("LINEからコピーしたシフトを入力してください。")
    rows, errors, seen = [], [], set()
    pattern = r"^(.+?)\s+(\d{4}[-/－／]\d{1,2}[-/－／]\d{1,2}|\d{1,2}[/／]\d{1,2})\s+(休み|\d{1,2}[:：]\d{2}\s*[-－〜～]\s*\d{1,2}[:：]\d{2})$"
    for number, line in enumerate(text.splitlines(), 1):
        if not line.strip():
            continue
        match = re.fullmatch(pattern, line.strip())
        if not match:
            errors.append(f"{number}行目：『山田 太郎 10/20 17:00-22:00』または『山田 太郎 10/20 休み』の形式で入力してください。")
            continue
        name, day_text, shift = match.groups()
        shift = normalize_numeric(shift)
        day = header_date(day_text, year, month)
        if not day or (day.year, day.month) != (year, month):
            errors.append(f"{number}行目：日付が存在しないか、対象月と一致しません。")
            continue
        if shift != "休み":
            start, end = re.split(r"\s*[-〜～]\s*", shift)
            parts = [tuple(map(int, t.split(":"))) for t in (start, end)]
            if any(not (0 <= h <= 23 and 0 <= m <= 59) for h, m in parts):
                errors.append(f"{number}行目：時刻は00:00〜23:59で入力してください。")
                continue
            if parts[1] <= parts[0]:
                errors.append(f"{number}行目：終了時刻は開始時刻より後にしてください。日をまたぐシフトは未対応です。")
                continue
            shift = "-".join(f"{h:02}:{m:02}" for h, m in parts)
        key = (normalize_name(name), day)
        if key in seen:
            errors.append(f"{number}行目：同じ人・同じ日のシフトが重複しています。")
            continue
        seen.add(key)
        rows.append({"name": name.strip(), "date": day.isoformat(), "shift": shift})
    if errors:
        raise InputError("\n".join(errors))
    if not rows:
        raise InputError("取り込めるシフトがありません。")
    return rows


def make_plan(book, data):
    year, month = month_parts(data.get("month", ""))
    sheet_name = data.get("sheet", "")
    if sheet_name not in book.sheetnames:
        raise InputError("反映先のシートを選んでください。")
    sheet = book[sheet_name]
    try:
        header = int(data.get("header_row", 1))
        first = int(data.get("first_name_row", 2))
        name_col = column_index_from_string(data.get("name_col", "A").strip().upper())
        date_col = column_index_from_string(data.get("date_col", "B").strip().upper())
        if not (1 <= header < first <= 10000 and name_col < date_col <= 16384):
            raise ValueError()
    except (ValueError, TypeError, AttributeError):
        raise InputError("表の位置を確認してください。氏名開始行は日付行より下、日付開始列は氏名列より右に指定してください。")
    if sheet.max_row > 10000 or sheet.max_column > 1000:
        raise InputError("表が大きすぎます（最大10000行・1000列）。不要な行列を削除したコピーを使用してください。")
    names, dates = {}, {}
    for row in range(first, sheet.max_row + 1):
        name = normalize_name(sheet.cell(row, name_col).value)
        if name:
            if name in names:
                raise InputError(f"Excelの氏名が重複しています：{sheet.cell(row, name_col).value}")
            names[name] = row
    for col in range(date_col, sheet.max_column + 1):
        day = header_date(sheet.cell(header, col).value, year, month)
        if day and (day.year, day.month) == (year, month):
            if day.isoformat() in dates:
                raise InputError(f"Excelの日付が重複しています：{day}")
            dates[day.isoformat()] = col
    result, errors = [], []
    for item in parse_lines(data.get("text", ""), data["month"]):
        row, col = names.get(normalize_name(item["name"])), dates.get(item["date"])
        if row is None or col is None:
            errors.append(f"{item['name']} {item['date']}：Excelに氏名または日付が見つかりません。")
            continue
        cell = sheet.cell(row, col)
        if sheet.protection.sheet:
            errors.append("シートが保護されています。保護を解除したコピーを使用してください。")
            break
        if any(cell.coordinate in merged for merged in sheet.merged_cells.ranges):
            errors.append(f"{cell.coordinate}は結合セルのため書き込めません。")
            continue
        if cell.data_type == "f":
            errors.append(f"{cell.coordinate}には数式が入っています。数式のない記入欄を指定してください。")
            continue
        existing = cell.value
        conflict = existing is not None and str(existing) != item["shift"]
        result.append({**item, "cell": cell.coordinate, "previous": str(existing) if existing is not None else "", "conflict": conflict})
    if errors:
        raise InputError("\n".join(errors))
    return sheet, result


def export_book(book, data):
    sheet, plan = make_plan(book, data)
    if any(item["conflict"] for item in plan) and data.get("overwrite") is not True:
        raise InputError("既に値があるセルがあります。確認表を確認して、上書き許可を選んでください。")
    for item in plan:
        sheet[item["cell"]] = item["shift"]
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def sample_book(month_value):
    year, month = month_parts(month_value)
    book = Workbook()
    sheet = book.active
    sheet.title = "シフト表"
    sheet.cell(1, 1, "氏名")
    for day in range(1, calendar.monthrange(year, month)[1] + 1):
        cell = sheet.cell(1, day + 1, date(year, month, day))
        cell.number_format = 'm/d'
        sheet.column_dimensions[cell.column_letter].width = 16
    for row, name in enumerate(["山田 太郎", "佐藤 花子", "鈴木 一郎"], 2):
        sheet.cell(row, 1, name)
    sheet.column_dimensions["A"].width = 20
    sheet.freeze_panes = "B2"
    output = BytesIO()
    book.save(output)
    return output.getvalue()

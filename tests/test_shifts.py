import base64
from io import BytesIO
import unittest

from openpyxl import load_workbook
from shifts import InputError, parse_lines, read_book, make_plan, export_book, sample_book


class ShiftTests(unittest.TestCase):
    def setUp(self):
        self.book = load_workbook(BytesIO(sample_book("2026-10")))
        self.data = {"month": "2026-10", "sheet": "シフト表", "text": "山田太郎 10/20 17:00-22:00\n佐藤 花子 10/21 休み", "header_row": 1, "first_name_row": 2, "name_col": "A", "date_col": "B"}

    def tearDown(self):
        self.book.close()

    def test_export_only_requested_cells_and_sheet(self):
        other = self.book.create_sheet("売上")
        other["A1"] = "=SUM(1,2)"
        self.book["シフト表"]["U2"].number_format = "@"
        result = load_workbook(BytesIO(export_book(self.book, self.data)))
        self.assertEqual(result["シフト表"]["U2"].value, "17:00-22:00")
        self.assertEqual(result["シフト表"]["V3"].value, "休み")
        self.assertIsNone(result["シフト表"]["T2"].value)
        self.assertEqual(result["シフト表"]["U2"].number_format, "@")
        self.assertEqual(result["売上"]["A1"].value, "=SUM(1,2)")
        result.close()

    def test_existing_value_needs_permission(self):
        self.book["シフト表"]["U2"] = "休み"
        _, rows = make_plan(self.book, self.data)
        self.assertTrue(rows[0]["conflict"])
        with self.assertRaises(InputError):
            export_book(self.book, self.data)
        self.assertEqual(self.book["シフト表"]["U2"].value, "休み")
        self.assertTrue(export_book(self.book, {**self.data, "overwrite": True}))

    def test_invalid_lines_dates_times_duplicates(self):
        for text in ["こんにちは", "山田太郎 10/32 休み", "山田太郎 11/1 休み", "山田太郎 10/20 25:00-26:00", "山田太郎 10/20 22:00-02:00", "山田太郎 10/20 休み\n山田 太郎 10/20 休み"]:
            with self.subTest(text=text), self.assertRaises(InputError):
                parse_lines(text, "2026-10")

    def test_unmatched_name_and_day_block_entire_export(self):
        for text in ["未登録 10/20 休み", "山田太郎 10/20 休み"]:
            data = {**self.data, "text": text}
            if text.startswith("山田"):
                self.book["シフト表"]["U1"] = None
            with self.assertRaises(InputError):
                export_book(self.book, data)
            self.assertIsNone(self.book["シフト表"]["V3"].value)

    def test_formula_merged_protected_cells_rejected(self):
        sheet = self.book["シフト表"]
        sheet["U2"] = "=1+1"
        with self.assertRaises(InputError):
            export_book(self.book, {**self.data, "overwrite": True})
        sheet["U2"] = None
        sheet.merge_cells("U2:V2")
        with self.assertRaises(InputError):
            export_book(self.book, self.data)
        sheet.unmerge_cells("U2:V2")
        sheet.protection.sheet = True
        with self.assertRaises(InputError):
            export_book(self.book, self.data)

    def test_custom_sheet_layout_and_numeric_headers(self):
        sheet = self.book.create_sheet("希望")
        sheet["B4"] = "山田 太郎"
        sheet["D3"] = 20
        data = {**self.data, "sheet": "希望", "name_col": "B", "date_col": "D", "header_row": 3, "first_name_row": 4, "text": "山田太郎 2026-10-20 9:00〜18:00"}
        _, rows = make_plan(self.book, data)
        self.assertEqual(rows[0]["cell"], "D4")
        self.assertEqual(rows[0]["shift"], "09:00-18:00")

    def test_duplicate_excel_headers_and_names_rejected(self):
        sheet = self.book["シフト表"]
        sheet["A5"] = "山田太郎"
        with self.assertRaises(InputError):
            make_plan(self.book, self.data)
        sheet["A5"] = None
        sheet["V1"] = sheet["U1"].value
        with self.assertRaises(InputError):
            make_plan(self.book, self.data)

    def test_file_import_and_invalid_upload(self):
        encoded = base64.b64encode(sample_book("2026-10")).decode()
        imported = read_book(encoded)
        self.assertEqual(imported.sheetnames, ["シフト表"])
        imported.close()
        with self.assertRaises(InputError):
            read_book("not an excel file")


if __name__ == "__main__":
    unittest.main()

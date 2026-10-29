import base64
from io import BytesIO
import tempfile
import unittest

from openpyxl import load_workbook

from app import create_app
from daily_shifts import sample_daily_book
from shifts import InputError, parse_lines, make_plan, sample_book


class FullWidthTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.app = create_app({"TESTING": True, "PUBLIC": False, "DATA_DIR": self.temp.name})
        self.client = self.app.test_client()

    def tearDown(self):
        self.temp.cleanup()

    def post(self, path, data):
        return self.client.post(path, json=data, headers={"Origin": "http://localhost"})

    def submit(self, shift="９：００-２１：００"):
        result = self.post("/api/submit", {"key": self.app.config["INVITE_CODE"], "name": "従業員１", "month": "２０２６－１１", "days": [{"date": "２０２６－１１－２０", "shift": shift}]})
        self.assertEqual(result.status_code, 200, result.json)

    def daily_data(self):
        self.post("/api/roster", {"kitchen": ["従業員１"], "hall": []})
        result = self.client.get("/api/daily-candidates?day=2026-11-20").json
        number = result["rows"][0]["id"]
        return {"file": base64.b64encode(sample_daily_book("2026-11-20")).decode(), "sheet": "日別ベータ", "day": "2026-11-20", "date_confirmed": True, "version": result["version"], "selected_ids": [number], "breaks": {str(number): {"mode": "time", "start": "１５：００", "end": "１７：００"}}}

    def test_pasted_fullwidth_date_time_preserves_employee_name(self):
        rows = parse_lines("従業員１ １１／２０ ９：００－２１：００\n従業員１ ２０２６／１１／２１ 休み", "２０２６－１１")
        self.assertEqual(rows, [{"name": "従業員１", "date": "2026-11-20", "shift": "09:00-21:00"}, {"name": "従業員１", "date": "2026-11-21", "shift": "休み"}])

    def test_submission_stored_under_normal_month_with_normal_times(self):
        self.submit()
        people = self.client.get("/api/submissions?month=2026-11").json
        self.assertEqual(people["people"][0]["name"], "従業員１")
        self.assertIn("2026-11-20 09:00-21:00", people["text"])
        self.assertEqual(self.client.get("/api/submissions?month=２０２６－１１").json, people)

    def test_manual_fullwidth_break_reaches_excel_as_time(self):
        self.submit(); data = self.daily_data()
        preview = self.post("/api/daily-preview", data)
        self.assertEqual(preview.status_code, 200, preview.json)
        self.assertEqual(preview.json["placements"][0]["break"], "15:00-17:00")
        result = self.post("/api/daily-export", {**data, "plan_token": preview.json["plan_token"]})
        self.assertEqual(result.status_code, 200)
        book = load_workbook(BytesIO(result.data))
        self.assertEqual(book.active["C15"].value, "従業員１")
        self.assertEqual(str(book.active["D15"].value), "9:00:00")
        self.assertEqual(str(book.active["H15"].value), "15:00:00")
        self.assertEqual(str(book.active["I15"].value), "17:00:00")
        book.close()

    def test_fixed_fullwidth_break_persisted_in_halfwidth(self):
        settings = {"mode": "auto", "weekday": {"mode": "time", "start": "１５：００", "end": "１７：００"}, "weekend": {"mode": "time", "start": "１６：００", "end": "１７：００"}}
        response = self.post("/api/break-settings", settings)
        self.assertEqual(response.status_code, 200, response.json)
        self.assertEqual(response.json["settings"]["weekday"]["start"], "15:00")
        self.client = create_app({"TESTING": True, "PUBLIC": False, "DATA_DIR": self.temp.name}).test_client()
        self.assertEqual(self.client.get("/api/break-settings").json["settings"]["weekend"]["start"], "16:00")

    def test_bad_fullwidth_times_not_repaired_into_valid_times(self):
        for shift in ["２５：００－２６：００", "１７：６０－２２：００", "２２：００－１７：００"]:
            with self.subTest(shift=shift), self.assertRaises(InputError):
                parse_lines(f"従業員１ １１／２０ {shift}", "2026-11")
        self.submit(); data = self.daily_data()
        for rest in data["breaks"].values(): rest["start"] = "１５：０１"
        self.assertIn("15分単位", self.post("/api/daily-preview", data).json["error"])

    def test_excel_header_accepts_fullwidth_date_text(self):
        book = load_workbook(BytesIO(sample_book("2026-11")))
        book.active["A2"] = "従業員１"
        book.active["U1"] = "１１／２０"
        _, rows = make_plan(book, {"month": "2026-11", "sheet": "シフト表", "text": "従業員１ １１／２０ １７：００－２２：００"})
        self.assertEqual(rows[0]["cell"], "U2")
        self.assertEqual(rows[0]["shift"], "17:00-22:00")
        book.close()


if __name__ == "__main__":
    unittest.main()

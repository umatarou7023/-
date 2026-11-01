import base64
from io import BytesIO
from pathlib import Path
import sqlite3
import tempfile
import unittest
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import load_workbook

from app import create_app
from daily_shifts import GROUPS, MAIN_NS, sample_daily_book


class DailyTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.app = create_app({"TESTING": True, "PUBLIC": False, "DATA_DIR": self.temp.name})
        self.client = self.app.test_client()
        self.source = sample_daily_book("2026-11-20")
        self.roster = {"kitchen": [f"厨房{i:02}" for i in range(1, 11)], "hall": [f"接客{i:02}" for i in range(1, 13)]}
        self.assertEqual(self.post("/api/roster", self.roster).status_code, 200)

    def tearDown(self):
        self.temp.cleanup()

    def post(self, url, data):
        return self.client.post(url, json=data, headers={"Origin": "http://localhost"})

    def submit(self, name, shift="09:00-21:00"):
        result = self.post("/api/submit", {"key": self.app.config["INVITE_CODE"], "name": name, "month": "2026-11", "days": [{"date": "2026-11-20", "shift": shift}]})
        self.assertEqual(result.status_code, 200, result.json)
        return result.json["receipt"]

    def data(self):
        candidates = self.client.get("/api/daily-candidates?day=2026-11-20").json
        ids = [row["id"] for row in candidates["rows"] if row["role"]]
        return {"file": base64.b64encode(self.source).decode(), "sheet": "日別ベータ", "day": "2026-11-20", "version": candidates["version"], "selected_ids": ids, "date_confirmed": True, "breaks": {str(i): {"mode": "none"} for i in ids}}

    def preview(self, data):
        result = self.post("/api/daily-preview", data)
        self.assertEqual(result.status_code, 200, result.json)
        return result.json

    def export(self, data, preview, overwrite=False):
        return self.post("/api/daily-export", {**data, "plan_token": preview["plan_token"], "overwrite": overwrite})

    def test_both_capacity_limits_block_without_losing_submissions(self):
        kitchen = [self.submit(n) for n in self.roster["kitchen"][:9]]
        hall = [self.submit(n) for n in self.roster["hall"][:11]]
        data = self.data()
        result = self.post("/api/daily-preview", data)
        self.assertEqual(result.status_code, 400)
        self.assertIn("キッチンは9人", result.json["error"])
        data["selected_ids"] = kitchen[:8] + hall
        self.assertIn("ホールは11人", self.post("/api/daily-preview", data).json["error"])
        data["selected_ids"] = kitchen[:8] + hall[:10]
        plan = self.preview(data)
        self.assertEqual(len(plan["placements"]), 18)
        self.assertEqual(len(plan["excluded"]), 2)
        self.assertEqual(self.export(data, plan).status_code, 200)
        with sqlite3.connect(Path(self.temp.name)/"shifts.sqlite3") as db:
            self.assertEqual(db.execute("SELECT COUNT(*) FROM submissions").fetchone()[0], 20)

    def test_submission_order_even_if_dinner_is_earlier_than_lunch(self):
        self.submit("厨房03", "17:00-21:00")
        self.submit("厨房02", "09:00-18:00")
        self.submit("厨房01", "09:00-18:00")
        people = self.preview(self.data())["placements"]
        self.assertEqual([p["name"] for p in people], ["厨房03", "厨房02", "厨房01"])
        self.assertEqual([p["cell"] for p in people], ["C15", "C17", "C19"])

    def test_resubmission_preserves_initial_priority_but_uses_latest_times(self):
        first = self.submit("厨房03", "17:00-21:00")
        self.submit("厨房01", "09:00-18:00")
        latest = self.submit("厨房03", "18:00-22:00")
        rows = self.client.get("/api/daily-candidates?day=2026-11-20").json["rows"]
        self.assertEqual(rows[0]["submission_order"], first)
        self.assertEqual(rows[0]["id"], latest)
        plan = self.preview(self.data())
        self.assertEqual([p["name"] for p in plan["placements"]], ["厨房03", "厨房01"])
        self.assertEqual(plan["placements"][0]["shift"], "18:00-22:00")

    def test_submission_priority_is_separate_for_each_month_and_role(self):
        self.submit("接客02", "17:00-21:00")
        self.submit("厨房02", "17:00-21:00")
        self.submit("接客01", "09:00-15:00")
        self.submit("厨房01", "09:00-15:00")
        plan = self.preview(self.data())["placements"]
        self.assertEqual([(p["name"], p["cell"]) for p in plan], [("厨房02", "C15"), ("厨房01", "C17"), ("接客02", "C34"), ("接客01", "C36")])
        for name in ["厨房01", "厨房02"]:
            result = self.post("/api/submit", {"key": self.app.config["INVITE_CODE"], "name": name, "month": "2026-12", "days": [{"date": "2026-12-20", "shift": "09:00-15:00"}]})
            self.assertEqual(result.status_code, 200)
        rows = self.client.get("/api/daily-candidates?day=2026-12-20").json["rows"]
        self.assertEqual([p["name"] for p in rows], ["厨房01", "厨房02"])

    def test_break_confirmed_and_written_as_excel_time(self):
        number = self.submit("厨房01")
        data = self.data()
        data["breaks"] = {}
        self.assertIn("休憩を確認", self.post("/api/daily-preview", data).json["error"])
        data["breaks"] = {str(number): {"mode": "time", "start": "15:00", "end": "17:00"}}
        plan = self.preview(data)
        result = self.export(data, plan)
        book = load_workbook(BytesIO(result.data))
        self.assertEqual(str(book.active["H15"].value), "15:00:00")
        self.assertEqual(str(book.active["I15"].value), "17:00:00")
        book.close()
        data["breaks"][str(number)]["end"] = "22:00"
        self.assertEqual(self.post("/api/daily-preview", data).status_code, 400)

    def test_existing_unused_slots_clear_only_after_explicit_permission(self):
        book = load_workbook(BytesIO(self.source))
        book.active["C19"] = "古い名前"
        book.active["D19"] = 9/24
        out = BytesIO(); book.save(out); book.close()
        self.source = out.getvalue()
        self.submit("厨房01")
        data = self.data(); plan = self.preview(data)
        self.assertTrue(any(u["cell"] == "C19" and u["next"] == "" and u["conflict"] for u in plan["updates"]))
        self.assertEqual(self.export(data, plan).status_code, 400)
        result = self.export(data, plan, overwrite=True)
        book = load_workbook(BytesIO(result.data))
        self.assertIsNone(book.active["C19"].value)
        self.assertIsNone(book.active["D19"].value)
        book.close()

    def test_xml_parts_formulas_and_conditional_formatting_preserved(self):
        self.submit("厨房01")
        data = self.data(); plan = self.preview(data)
        result = self.export(data, plan)
        self.assertEqual(result.status_code, 200)
        with ZipFile(BytesIO(self.source)) as before, ZipFile(BytesIO(result.data)) as after:
            self.assertEqual(before.namelist(), after.namelist())
            for name in before.namelist():
                if name not in {"xl/worksheets/sheet1.xml", "xl/workbook.xml"}:
                    self.assertEqual(before.read(name), after.read(name), name)
            old, new = [ET.fromstring(z.read("xl/worksheets/sheet1.xml")) for z in (before, after)]
            for tag in ("f", "conditionalFormatting", "mergeCells"):
                self.assertEqual([ET.tostring(n) for n in old.iter(f"{{{MAIN_NS}}}{tag}")], [ET.tostring(n) for n in new.iter(f"{{{MAIN_NS}}}{tag}")])

    def test_changed_submissions_or_roster_require_reload(self):
        self.submit("厨房01")
        data = self.data(); plan = self.preview(data)
        self.submit("厨房01", "休み")
        self.assertEqual(self.export(data, plan).status_code, 400)
        self.assertIn("読み込み直", self.export(data, plan).json["error"])
        self.submit("厨房01")
        data = self.data(); plan = self.preview(data)
        self.post("/api/roster", {**self.roster, "kitchen": list(reversed(self.roster["kitchen"]))})
        self.assertEqual(self.export(data, plan).status_code, 400)

    def test_changed_plan_cannot_skip_review(self):
        number = self.submit("厨房01")
        data = self.data(); plan = self.preview(data)
        data["breaks"][str(number)] = {"mode": "time", "start": "12:00", "end": "13:00"}
        self.assertIn("プレビュー", self.export(data, plan).json["error"])

    def test_unregistered_names_and_days_off_remain_saved(self):
        unknown = self.submit("未登録")
        self.submit("厨房01", "休み")
        data = self.data()
        result = self.client.get("/api/daily-candidates?day=2026-11-20").json
        self.assertEqual(len(result["rows"]), 1)
        self.assertTrue(result["rows"][0]["issue"])
        data["selected_ids"] = [unknown]
        self.assertIn("未登録", self.post("/api/daily-preview", data).json["error"])
        self.assertEqual(len(self.client.get("/api/submissions?month=2026-11").json["people"]), 2)

    def test_date_template_formula_and_rounding_mismatch_block(self):
        self.submit("厨房01")
        data = self.data(); data["date_confirmed"] = False
        self.assertEqual(self.post("/api/daily-preview", data).status_code, 400)
        for coordinate, value in [("B9", 21), ("D15", "=1"), ("C14", "移動した名前欄")]:
            book = load_workbook(BytesIO(self.source)); book.active[coordinate] = value
            out = BytesIO(); book.save(out); book.close()
            data = self.data(); data["file"] = base64.b64encode(out.getvalue()).decode()
            self.assertEqual(self.post("/api/daily-preview", data).status_code, 400, coordinate)
        self.submit("厨房01", "09:01-18:00")
        self.assertIn("15分単位", self.post("/api/daily-preview", self.data()).json["error"])
        self.submit("厨房01", "08:00-18:00")
        self.assertIn("表示範囲外", self.post("/api/daily-preview", self.data()).json["error"])

    def test_duplicate_roster_and_formula_like_names(self):
        self.assertEqual(self.post("/api/roster", {"kitchen": ["山田 太郎"], "hall": ["山田太郎"]}).status_code, 400)
        self.post("/api/roster", {"kitchen": ["=1+1"], "hall": []})
        self.submit("=1+1")
        data = self.data(); result = self.export(data, self.preview(data))
        book = load_workbook(BytesIO(result.data))
        self.assertEqual(book.active["C15"].data_type, "s")
        self.assertEqual(book.active["C15"].value, "=1+1")
        book.close()


if __name__ == "__main__":
    unittest.main()

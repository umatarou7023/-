"""Book3型の日別シフト。計算欄を変更せず入力セルだけを更新する。"""
import base64
from datetime import date, time, timedelta
import hashlib
from io import BytesIO
import json
import re
from xml.dom import minidom
import xml.etree.ElementTree as ET
from zipfile import ZipFile

from shifts import InputError, normalize_name, normalize_numeric, normalize_clock

GROUPS = {"kitchen": ("キッチン", list(range(15, 31, 2))),
          "hall": ("ホール", list(range(34, 54, 2)))}
MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


def fingerprint(value):
    return hashlib.sha256(json.dumps(value, ensure_ascii=False, sort_keys=True).encode()).hexdigest()


def parse_roster(value):
    if not isinstance(value, dict):
        raise InputError("スタッフ一覧を登録してください。")
    result, seen = {}, set()
    for role, (label, _) in GROUPS.items():
        names = value.get(role, [])
        if not isinstance(names, list) or len(names) > 300:
            raise InputError(f"{label}の一覧は300人以内で登録してください。")
        result[role] = []
        for name in names:
            if not isinstance(name, str) or not name.strip() or len(name) > 60 or len(name.splitlines()) != 1:
                raise InputError("スタッフ一覧は1行に1人の氏名を入力してください。")
            key = normalize_name(name)
            if key in seen:
                raise InputError(f"スタッフ一覧で氏名が重複しています：{name}")
            seen.add(key)
            result[role].append(name.strip())
    return result


def target_day(value):
    try:
        value = normalize_numeric(value)
        if not isinstance(value, str) or not re.fullmatch(r"\d{4}-\d{2}-\d{2}", value):
            raise ValueError()
        return date.fromisoformat(value)
    except ValueError:
        raise InputError("反映する日付を選んでください。")


def daily_candidates(records, roster, day):
    target_day(day)
    members = {normalize_name(name): (role, order, name)
               for role, names in roster.items() for order, name in enumerate(names)}
    result = []
    for record in records:
        for row in record["rows"]:
            if row["date"] != day or row["shift"] == "休み":
                continue
            member = members.get(normalize_name(row["name"]))
            result.append({"id": record["id"], "name": member[2] if member else row["name"],
                           "role": member[0] if member else "", "order": member[1] if member else 0,
                           "shift": row["shift"], "issue": "" if member else "スタッフ一覧に未登録"})
    # 休みへの再提出や、その日の希望の削除も、読み込み済み画面を無効にする。
    return result, fingerprint({"day": day, "records": records, "roster": roster})


def validate_template(book, sheet_name):
    if sheet_name not in book.sheetnames:
        raise InputError("反映先シートを選んでください。")
    sheet = book[sheet_name]
    if sheet.protection.sheet:
        raise InputError("シートが保護されています。元ファイルの管理者に確認してください。")
    if sheet["B15"].value != "キッチン" or sheet["B34"].value not in {"フロア", "ホール"}:
        raise InputError("キッチン・ホールの配置が確認済みの形式と違います。")
    for row in (14, 33):
        for col, text in {"C": "名前", "D": "開始時間", "E": "終了時間", "H": "休憩開始", "I": "休憩終了"}.items():
            if sheet[f"{col}{row}"].value != text:
                raise InputError("確認済みの日別表と配置が違います。推測で書き込まず停止しました。")
    for _, slots in GROUPS.values():
        for row in slots:
            for col in ("C", "D", "E", "H", "I"):
                if f"{col}{row}:{col}{row+1}" not in {str(r) for r in sheet.merged_cells.ranges}:
                    raise InputError(f"{col}{row}の結合範囲が確認済みの形式と違います。")
                if sheet[f"{col}{row}"].data_type == "f":
                    raise InputError(f"入力先{col}{row}に数式があるため停止しました。")
            expected = {dest: f'={rounding}({source}{row},"0:15")'
                        for dest, source, rounding in (("F", "D", "CEILING"), ("G", "E", "FLOOR"), ("J", "H", "CEILING"), ("K", "I", "FLOOR"))}
            expected["L"] = f'=IF(AND($J{row}<=L$56,$K{row}>L$56),"・・",IF(AND($F{row}<=L$56,$G{row}>L$56),"・",""))'
            for col, formula in expected.items():
                if sheet[f"{col}{row}"].value != formula:
                    raise InputError(f"{col}{row}の計算式が確認済みの形式と違います。表の構成を確認してください。")
    return sheet


def minutes(value):
    value = normalize_clock(value)
    if not isinstance(value, str) or not re.fullmatch(r"\d{2}:\d{2}", value):
        raise InputError("時刻はHH:MM形式で指定してください。")
    hour, minute = map(int, value.split(":"))
    if not (0 <= hour <= 23 and 0 <= minute <= 59) or minute % 15:
        raise InputError("この日別表のベータ版は15分単位の時刻に対応しています。端数を確認してください。")
    return hour * 60 + minute


def show_value(value):
    if value is None:
        return ""
    if isinstance(value, timedelta):
        total = round(value.total_seconds() / 60)
        return f"{total // 60:02}:{total % 60:02}"
    if isinstance(value, time):
        return value.strftime("%H:%M")
    if isinstance(value, (int, float)) and 0 <= value < 1:
        total = round(value * 1440)
        return f"{total // 60:02}:{total % 60:02}"
    return str(value)


def make_daily_plan(book, data, candidates):
    sheet = validate_template(book, data.get("sheet", ""))
    day = target_day(data.get("day"))
    try:
        axis_start = minutes(show_value(sheet["L5"].value))
    except InputError:
        raise InputError("勤務帯の表示開始時刻（L5）が確認できません。Excel側の設定を確認してください。")
    axis_end = axis_start + 64 * 15  # L〜BWの64枠。1枠15分。
    # 日付の式は触らない。既知の日付と違う場合は明示確認でも通さない。
    if sheet["E5"].value != "=DATE(C4,B5,B9)":
        raise InputError("日付の計算方法が確認済みの形式と違います。")
    try:
        workbook_day = date(*(int(sheet[cell].value) for cell in ("C4", "B5", "B9")))
    except (ValueError, TypeError):
        workbook_day = None
    if workbook_day and workbook_day != day:
        raise InputError(f"Excelの日付（{workbook_day}）と反映日（{day}）が違います。Excelのコピー側の日付を設定してください。")
    if data.get("date_confirmed") is not True:
        raise InputError("このシートが対象日の表であることを確認してください。日付未設定のExcelは出力後に設定が必要です。")
    ids = data.get("selected_ids")
    if not isinstance(ids, list) or not ids or any(type(i) is not int for i in ids) or len(ids) != len(set(ids)):
        raise InputError("反映する人を1人以上選択してください。")
    by_id = {c["id"]: c for c in candidates}
    if any(i not in by_id for i in ids):
        raise InputError("提出内容が変わりました。一覧を読み込み直してください。")
    chosen = [by_id[i] for i in ids]
    if any(c["issue"] for c in chosen):
        raise InputError("未登録の従業員がいます。スタッフ一覧を確認してから読み込み直してください。")
    for role, (label, slots) in GROUPS.items():
        count = sum(c["role"] == role for c in chosen)
        if count > len(slots):
            raise InputError(f"{label}は{count}人選択中で、{count-len(slots)}人分超えています（上限{len(slots)}人）。管理者が対象者を選び直してください。提出内容は削除されません。")
    breaks = data.get("breaks", {})
    if not isinstance(breaks, dict):
        raise InputError("休憩を確認してください。")
    updates, placements = [], []
    for role, (label, slots) in GROUPS.items():
        people = sorted((c for c in chosen if c["role"] == role), key=lambda c: (minutes(c["shift"].split("-")[0]), c["order"]))
        for index, row in enumerate(slots):
            person = people[index] if index < len(people) else None
            values = {"C": None, "D": None, "E": None, "H": None, "I": None}
            if person:
                start_text, end_text = person["shift"].split("-")
                start, end = minutes(start_text), minutes(end_text)
                if end <= start:
                    raise InputError("終了時刻は開始時刻より後にしてください。")
                if not (axis_start <= start < end <= axis_end):
                    raise InputError(f"{person['name']}の勤務時間が、このExcelの表示範囲外です。時間帯の設定を確認してください。")
                rest = breaks.get(str(person["id"]), {})
                if not isinstance(rest, dict) or rest.get("mode") not in {"none", "time"}:
                    raise InputError(f"{person['name']}の休憩を確認してください（休憩なし／時刻指定）。")
                rest_text = "なし（管理者確認）"
                if rest["mode"] == "time":
                    rest_start, rest_end = normalize_clock(rest.get("start")), normalize_clock(rest.get("end"))
                    rs, re_ = minutes(rest_start), minutes(rest_end)
                    if not (start <= rs < re_ <= end):
                        raise InputError(f"{person['name']}の休憩は勤務時間内で、開始より終了を後にしてください。")
                    values["H"], values["I"] = rs / 1440, re_ / 1440
                    rest_text = f"{rest_start}-{rest_end}"
                values.update(C=person["name"], D=start/1440, E=end/1440)
                placements.append({"id": person["id"], "role": label, "name": person["name"], "shift": person["shift"], "break": rest_text, "cell": f"C{row}", "date": day.isoformat()})
            for col, value in values.items():
                cell = sheet[f"{col}{row}"]
                previous, new = show_value(cell.value), (show_value(value) if col != "C" else (value or ""))
                # 0:00の値を文字列化してから比較し、空欄と区別する。
                if previous == new:
                    continue
                updates.append({"cell": cell.coordinate, "value": value, "previous": previous, "next": new,
                                "conflict": cell.value is not None, "name": person["name"] if person else "未使用枠を空欄に戻す"})
    excluded = [c["name"] for c in candidates if c["id"] not in ids]
    return {"placements": placements, "updates": updates, "excluded": excluded,
            "date_missing": workbook_day is None,
            "warnings": ["未使用の枠は名前・勤務時間・休憩時間を空欄に戻します。", "休憩は管理者の指定内容です。法令や職場のルールへの適合を自動判定する機能はありません。", "勤務帯・合計はExcelで再計算して確認してください。"]}


def patch_daily_excel(encoded, sheet_name, plan):
    """元のZIP各パーツを保持し、確認済み入力セルと再計算フラグだけを更新する。"""
    source = BytesIO(base64.b64decode(encoded, validate=True))
    result = BytesIO()
    with ZipFile(source) as zin:
        workbook = ET.fromstring(zin.read("xl/workbook.xml"))
        relation_id = next(s.attrib["{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"] for s in workbook.find(f"{{{MAIN_NS}}}sheets") if s.attrib["name"] == sheet_name)
        relationships = ET.fromstring(zin.read("xl/_rels/workbook.xml.rels"))
        target = next(r.attrib["Target"] for r in relationships if r.attrib["Id"] == relation_id)
        import posixpath
        sheet_path = target.lstrip("/") if target.startswith("/") else posixpath.normpath("xl/" + target)
        doc = minidom.parseString(zin.read(sheet_path))
        cells = {c.getAttribute("r"): c for c in doc.getElementsByTagNameNS(MAIN_NS, "c")}
        for item in plan["updates"]:
            coordinate, value = item["cell"], item["value"]
            cell = cells.get(coordinate)
            if cell is None:
                row_number = re.search(r"\d+$", coordinate).group()
                row = next(r for r in doc.getElementsByTagNameNS(MAIN_NS, "row") if r.getAttribute("r") == row_number)
                cell = doc.createElementNS(MAIN_NS, "c")
                cell.setAttribute("r", coordinate)
                # XML内のセル順を保つ。
                from openpyxl.utils import column_index_from_string
                position = column_index_from_string(re.match(r"[A-Z]+", coordinate).group())
                following = next((c for c in row.childNodes if c.nodeType == c.ELEMENT_NODE and c.localName == "c" and column_index_from_string(re.match(r"[A-Z]+", c.getAttribute("r")).group()) > position), None)
                row.insertBefore(cell, following)
            for child in list(cell.childNodes):
                cell.removeChild(child)
            if cell.hasAttribute("t"):
                cell.removeAttribute("t")
            if isinstance(value, str):
                cell.setAttribute("t", "inlineStr")
                inline = doc.createElementNS(MAIN_NS, "is")
                text = doc.createElementNS(MAIN_NS, "t")
                text.appendChild(doc.createTextNode(value))
                inline.appendChild(text)
                cell.appendChild(inline)
            elif value is not None:
                node = doc.createElementNS(MAIN_NS, "v")
                node.appendChild(doc.createTextNode(str(value)))
                cell.appendChild(node)
        wb_doc = minidom.parseString(zin.read("xl/workbook.xml"))
        calc_nodes = wb_doc.getElementsByTagNameNS(MAIN_NS, "calcPr")
        if calc_nodes:
            calc = calc_nodes[0]
        else:
            calc = wb_doc.createElementNS(MAIN_NS, "calcPr")
            wb_doc.documentElement.appendChild(calc)
        for key, value in {"calcMode": "auto", "fullCalcOnLoad": "1", "forceFullCalc": "1"}.items():
            calc.setAttribute(key, value)
        with ZipFile(result, "w") as zout:
            for info in zin.infolist():
                content = doc.toxml(encoding="utf-8") if info.filename == sheet_path else wb_doc.toxml(encoding="utf-8") if info.filename == "xl/workbook.xml" else zin.read(info.filename)
                zout.writestr(info, content)
    return result.getvalue()


def sample_daily_book(day_value):
    """実際の職場のファイル・従業員名を含まない、独立したベータ用の表。"""
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.utils import get_column_letter

    day = target_day(day_value)
    book = Workbook()
    sheet = book.active
    sheet.title = "日別ベータ"
    sheet["B2"] = "ベータテスト用・架空の勤務表"
    sheet["C4"], sheet["B5"], sheet["B9"] = day.year, day.month, day.day
    sheet["E5"] = "=DATE(C4,B5,B9)"
    sheet["E5"].number_format = "yyyy/m/d"
    sheet["B15"], sheet["B34"] = "キッチン", "ホール"
    sheet["L5"] = 9/24
    sheet["L5"].number_format = "h:mm"
    for header in (14, 33):
        for col, label in {"C": "名前", "D": "開始時間", "E": "終了時間", "H": "休憩開始", "I": "休憩終了"}.items():
            sheet[f"{col}{header}"] = label
    for col in range(12, 76):
        letter = get_column_letter(col)
        sheet[f"{letter}56"] = ((9*60+(col-12)*15) % 1440 + 1/60) / 1440
        for header in (14, 33):
            sheet[f"{letter}{header}"] = ((9*60+(col-12)*15) % 1440) / 1440
            sheet[f"{letter}{header}"].number_format = "h:mm"
        sheet.column_dimensions[letter].width = 3
    for _, slots in GROUPS.values():
        for row in slots:
            for col in ("C", "D", "E", "H", "I"):
                sheet.merge_cells(f"{col}{row}:{col}{row+1}")
                sheet[f"{col}{row}"].number_format = "@" if col == "C" else "[h]:mm"
            for dest, source, rounding in (("F", "D", "CEILING"), ("G", "E", "FLOOR"), ("J", "H", "CEILING"), ("K", "I", "FLOOR")):
                sheet[f"{dest}{row}"] = f'={rounding}({source}{row},"0:15")'
            for col in range(12, 76):
                letter = get_column_letter(col)
                sheet[f"{letter}{row}"] = f'=IF(AND($J{row}<={letter}$56,$K{row}>{letter}$56),"・・",IF(AND($F{row}<={letter}$56,$G{row}>{letter}$56),"・",""))'
                sheet[f"{letter}{row}"].number_format = ";;;"
            sheet[f"BX{row}"] = f'=COUNTIF(L{row}:BW{row},"・")/4'
            sheet.conditional_formatting.add(f"L{row}:BW{row}", CellIsRule(operator="equal", formula=['"・"'], fill=PatternFill("solid", fgColor="000000")))
    for col in ("C", "D", "E", "H", "I"):
        sheet.column_dimensions[col].width = 16
    for col in ("F", "G", "J", "K"):
        sheet.column_dimensions[col].hidden = True
    sheet.row_dimensions[56].hidden = True
    sheet.freeze_panes = "L15"
    sheet.print_area = "B1:BX53"
    out = BytesIO()
    book.save(out)
    return out.getvalue()
